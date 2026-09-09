# -*- coding: utf-8 -*-
"""Word 导出：将纪要/周报生成 .docx"""
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


ACCENT = RGBColor(0x2F, 0x6F, 0xED)
GRAY = RGBColor(0x6B, 0x72, 0x80)


def _style_base(doc):
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)
    # 中文字体兼容
    from docx.oxml.ns import qn
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _h(doc, text, size=14, color=ACCENT, space_before=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def _para(doc, text, bold=False, color=None, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def export_meeting_docx(meeting: dict) -> bytes:
    doc = Document()
    _style_base(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(meeting.get("title") or "会议纪要")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
                      f"　|　时长/消息量：{meeting.get('duration', '')}"
                      f"　|　参与人：{', '.join(meeting.get('participants', []) or ['-'])}")
    mr.font.size = Pt(9)
    mr.font.color.rgb = GRAY

    _h(doc, "摘要")
    _para(doc, meeting.get("summary", ""))

    _h(doc, f"议题讨论（{len(meeting.get('topics', []) or [])}）")
    for i, tp in enumerate(meeting.get("topics", []) or [], 1):
        _h(doc, f"{i}. {tp.get('title', '')}", size=12, color=RGBColor(0x1A, 0x1A, 0x1A))
        for pt in tp.get("points", []) or []:
            _para(doc, f"• {pt}")
        if tp.get("conclusion"):
            _para(doc, f"结论：{tp['conclusion']}", color=RGBColor(0x2F, 0x6F, 0xED))

    todos = meeting.get("todos", []) or []
    _h(doc, f"待办事项（{len(todos)}）")
    if todos:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, htext in enumerate(["任务", "负责人", "截止时间", "优先级"]):
            hdr[i].paragraphs[0].add_run(htext).bold = True
        for t in todos:
            row = table.add_row().cells
            row[0].text = str(t.get("task", ""))
            row[1].text = str(t.get("owner", "待确认"))
            row[2].text = str(t.get("deadline", "未定"))
            row[3].text = str(t.get("priority", "中"))

    risks = meeting.get("risks", []) or []
    if risks:
        _h(doc, "风险提醒")
        for rk in risks:
            _para(doc, f"• {rk}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_weekly_docx(weekly_text: str, agg) -> bytes:
    doc = Document()
    _style_base(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("团队周报")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"统计周期：本周　|　生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　|　"
                      f"会议场次：{agg['meeting_count']}　|　待办：{len(agg['todos'])}")
    mr.font.size = Pt(9)
    mr.font.color.rgb = GRAY

    # 统计表
    _h(doc, "任务分布（按负责人）")
    if not agg["owner_df"].empty:
        t1 = doc.add_table(rows=1, cols=len(agg["owner_df"].columns))
        t1.style = "Light Grid Accent 1"
        for i, c in enumerate(agg["owner_df"].columns):
            t1.rows[0].cells[i].paragraphs[0].add_run(str(c)).bold = True
        for _, row in agg["owner_df"].iterrows():
            cells = t1.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
    else:
        _para(doc, "本周暂无待办任务统计。")

    _h(doc, "周报正文")
    for line in weekly_text.split("\n"):
        if line.startswith("【"):
            _h(doc, line.strip("【】"), size=12, color=RGBColor(0x1A, 0x1A, 0x1A))
        elif line.strip():
            _para(doc, line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
